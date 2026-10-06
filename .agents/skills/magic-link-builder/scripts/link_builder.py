import os
import sys
import re
import argparse
import openpyxl

# Ajout du module partagé au path
_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_PROJECT_ROOT = os.path.abspath(os.path.join(_SCRIPT_DIR, '..', '..', '..', '..'))
if _PROJECT_ROOT not in sys.path:
    sys.path.insert(0, _PROJECT_ROOT)

from shared.config import (
    ensure_utf8, get_book_config, normalize_book_key,
    find_data_sheet, backup_excel, BOOKS_CONFIG, SOURCES_LISIBLES_DIR, BASE_DIR
)

ensure_utf8()


class MagicLinkBuilder:
    def __init__(self, sources_dir: str = SOURCES_LISIBLES_DIR):
        self.sources_dir = sources_dir
        self.pdf_index = self._index_sources()

    def _index_sources(self) -> dict:
        """Indexe tous les fichiers PDF certifiés OCR dans sources_lisibles."""
        index = {}
        if not os.path.exists(self.sources_dir):
            return index

        for root, dirs, files in os.walk(self.sources_dir):
            for f in files:
                if f.lower().endswith('.pdf'):
                    full_path = os.path.join(root, f)
                    rel_path = os.path.relpath(full_path, BASE_DIR).replace('\\', '/')
                    clean_key = f.strip().lower()
                    index[clean_key] = rel_path
        return index

    def resolve_single_link(self, ref_text: str):
        """
        Extrait le PDF et la page d'une référence unique.
        Retourne (target_path, page_num, link_url)
        """
        if not ref_text or not str(ref_text).strip():
            return None, None, None

        first_ref = str(ref_text).strip()

        # 1. Extraction du numéro de page
        page = 1
        m_sect = re.search(r'\(sect\.\s*~?(\d+)\)', first_ref, re.IGNORECASE)
        if m_sect:
            page = int(m_sect.group(1))
        else:
            m_page = re.search(r'[,\s]+p(?:age)?\.?\s*(\d+)', first_ref, re.IGNORECASE)
            if m_page:
                page = int(m_page.group(1))

        # 2. Correspondance exacte du nom de fichier
        first_ref_lower = first_ref.lower()
        for pdf_name, target_path in self.pdf_index.items():
            if pdf_name in first_ref_lower:
                return target_path, page, f"{target_path}#page={page}"

        # 3. Correspondance tolérante sur le tronc du nom
        for pdf_name, target_path in self.pdf_index.items():
            core_part = pdf_name.replace('.pdf', '')
            if len(core_part) > 10 and core_part[:25] in first_ref_lower:
                return target_path, page, f"{target_path}#page={page}"

        return None, None, None

    def process_workbook(self, book_key: str, dry_run: bool = False) -> int:
        config = get_book_config(book_key)
        excel_path = config['excel']

        if not os.path.exists(excel_path):
            print(f"❌ Fichier introuvable: {excel_path}")
            return 0

        print(f"\n=======================================================")
        print(f"🔗 SKILL AGENT 3 : magic-link-builder")
        print(f"=======================================================")
        print(f"📖 Livre cible     : {config['title']} ({config['id']})")
        print(f"📊 Fichier Excel   : {os.path.basename(excel_path)}")
        print(f"🔒 Mode simulation : {'OUI (Dry Run)' if dry_run else 'NON (Écriture Excel)'}\n")

        wb = openpyxl.load_workbook(excel_path)
        sheet = find_data_sheet(wb)

        total_rows = sheet.max_row - 1
        links_built = 0
        links_unresolved = 0

        for r in range(2, sheet.max_row + 1):
            ref_val = str(sheet.cell(r, 7).value or '').strip()
            
            if not ref_val:
                link_cell = sheet.cell(r, 8)
                if not dry_run:
                    link_cell.value = "PDF non indexé"
                    link_cell.hyperlink = None
                links_unresolved += 1
                continue
                
            refs = [rf.strip() for rf in ref_val.split('|') if rf.strip()]
            link_urls = []
            pages = []
            
            for rf in refs:
                target_path, page, link_url = self.resolve_single_link(rf)
                if link_url:
                    link_urls.append(link_url)
                    pages.append(str(page))

            link_cell = sheet.cell(r, 8)
            if link_urls:
                if not dry_run:
                    if len(link_urls) == 1:
                        link_cell.value = f"Ouvrir page {pages[0]}"
                        link_cell.hyperlink = link_urls[0]
                    else:
                        link_cell.value = " | ".join(link_urls)
                        link_cell.hyperlink = None
                links_built += 1
            else:
                if not dry_run:
                    link_cell.value = "PDF non indexé"
                    link_cell.hyperlink = None
                links_unresolved += 1

        if not dry_run and links_built > 0:
            backup_excel(excel_path)
            wb.save(excel_path)
            print(f"💾 Sauvegarde réussie dans {os.path.basename(excel_path)} ({links_built} / {total_rows} liens résolus).")
        else:
            print(f"🏁 Fin de l'opération ({links_built} / {total_rows} liens vérifiés).")

        return links_built

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description="Skill Agent 3 : magic-link-builder")
    parser.add_argument('--book', default='cartes', help="Clé du livre ('cartes', 'pieces', 'tours')")
    parser.add_argument('--all', action='store_true', help="Traiter les 3 livres")
    parser.add_argument('--dry-run', action='store_true', help="Mode simulation sans écriture")
    args = parser.parse_args()

    builder = MagicLinkBuilder()
    print(f"🔍 {len(builder.pdf_index)} fichiers PDF indexés depuis Sources/sources_lisibles/")

    if args.all:
        for bk in BOOKS_CONFIG.keys():
            builder.process_workbook(bk, dry_run=args.dry_run)
    else:
        builder.process_workbook(args.book, dry_run=args.dry_run)
