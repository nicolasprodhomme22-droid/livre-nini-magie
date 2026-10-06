import openpyxl
from openpyxl.worksheet.hyperlink import Hyperlink
import os
import sys
import re
import unicodedata
from collections import defaultdict
import time
import argparse

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_PROJECT_ROOT = os.path.abspath(os.path.join(_SCRIPT_DIR, '..', '..', '..', '..'))
if _PROJECT_ROOT not in sys.path:
    sys.path.insert(0, _PROJECT_ROOT)

from shared.config import ensure_utf8, find_data_sheet, backup_excel, BOOKS_CONFIG, BASE_DIR

ensure_utf8()

GENERIC_NAMES = {
    'move', 'take', 'palm', 'the palm', 'the coins', 'change', 'production',
    'set up', 'the end', 'the blind shuffle', 'i deal', 'dealing', 'routine'
}

def strip_accents(text):
    if not text:
        return ""
    return ''.join(c for c in unicodedata.normalize('NFD', text) if unicodedata.category(c) != 'Mn')

def deep_norm(text):
    if not text:
        return ""
    text = strip_accents(str(text)).lower()
    text = re.sub(r'[\(\)\[\]\.,;:!\?\"\'\-_]', ' ', text)
    text = re.sub(r'\s+', ' ', text).strip()
    return text

def parse_page_number(ref_str, target_str=""):
    combined = f"{ref_str} {target_str}"
    m = re.search(r'#page=(\d+)', combined, re.IGNORECASE)
    if m:
        return int(m.group(1))
    m2 = re.search(r'p\.\s*(\d+)', combined, re.IGNORECASE)
    if m2:
        return int(m2.group(1))
    m3 = re.search(r'page\s*(\d+)', combined, re.IGNORECASE)
    if m3:
        return int(m3.group(1))
    return 1

def parse_pdf_filename(ref_str, target_str=""):
    for s in [target_str, ref_str]:
        if not s:
            continue
        clean = s.split('#')[0].replace('\\', '/').strip()
        base = clean.split('/')[-1].split(',')[0].strip()
        if base.lower().endswith('.pdf'):
            return base
    parts = ref_str.split(',')
    if parts:
        cand = parts[0].strip()
        if cand:
            return cand
    return "document.pdf"

def deduplicate_workbook_fast(book_title, file_name):
    t0 = time.time()
    file_path = os.path.join(BASE_DIR, file_name)
    print("=" * 60)
    print(f"🧹 CONSOLIDATION ET DÉDUPLICATION : {book_title}")
    print(f"📄 Fichier : {file_name}")
    print("=" * 60)
    
    wb = openpyxl.load_workbook(file_path, data_only=False)
    sheet = find_data_sheet(wb)
        
    print(f"📑 Feuille active : {sheet.title} (Lignes : {sheet.max_row})")
    
    # 1. Lire toutes les lignes
    headers = [sheet.cell(1, c).value for c in range(1, 9)]
    
    rows_data = []
    for r in range(2, sheet.max_row + 1):
        id_val = sheet.cell(r, 1).value
        if not id_val:
            continue
        p = str(sheet.cell(r, 2).value or '').strip()
        c = str(sheet.cell(r, 3).value or '').strip()
        s = str(sheet.cell(r, 4).value or '').strip()
        nom = str(sheet.cell(r, 5).value or '').strip()
        desc = str(sheet.cell(r, 6).value or '').strip()
        pdf_ref = str(sheet.cell(r, 7).value or '').strip()
        
        cell_link = sheet.cell(r, 8)
        target = ""
        if cell_link.hyperlink and cell_link.hyperlink.target:
            target = cell_link.hyperlink.target
        elif cell_link.value and '#' in str(cell_link.value):
            target = str(cell_link.value).strip()
            
        norm = deep_norm(nom)
        if not norm:
            continue
            
        rows_data.append({
            'id': id_val,
            'partie': p,
            'chapitre': c,
            'section': s,
            'nom': nom,
            'desc': desc,
            'pdf_ref': pdf_ref,
            'target': target,
            'norm': norm
        })
        
    # Groupement
    groups = defaultdict(list)
    ordered_keys = []
    for r in rows_data:
        norm = r['norm']
        if norm in GENERIC_NAMES or len(norm) <= 3:
            key = (r['partie'], r['chapitre'], r['section'], norm)
        else:
            key = (r['partie'], norm)
        if key not in groups:
            ordered_keys.append(key)
        groups[key].append(r)
        
    dup_count = sum(1 for k in ordered_keys if len(groups[k]) > 1)
    print(f"🔍 {dup_count} groupes avec doublons identifiés sur {len(ordered_keys)} techniques distinctes.")
    
    # 2. Consolider chaque groupe en une SEULE ligne
    final_rows = []
    for key in ordered_keys:
        rlist = groups[key]
        if len(rlist) == 1:
            item = rlist[0]
            final_rows.append({
                'id': item['id'],
                'partie': item['partie'],
                'chapitre': item['chapitre'],
                'section': item['section'],
                'nom': item['nom'],
                'desc': item['desc'],
                'pdf_ref': item['pdf_ref'],
                'target': item['target'],
                'primary_target': item['target'].split('|')[0].strip() if item['target'] else ''
            })
        else:
            # Meilleure ligne primaire
            def row_score(item):
                score = 100 if item['section'] and item['section'] != 'Général' else 0
                if '[Méthode]' in item['desc'] and '[Effet]' in item['desc']:
                    score += 50 + min(len(item['desc']), 500)
                else:
                    score += min(len(item['desc']), 200)
                return score
                
            rlist_sorted = sorted(rlist, key=row_score, reverse=True)
            primary = rlist_sorted[0]
            others = rlist_sorted[1:]
            
            # Meilleur titre
            best_title = primary['nom']
            for o in others:
                if '(' in o['nom'] and ')' in o['nom'] and '(' not in best_title:
                    best_title = o['nom']
                elif len(o['nom']) > len(best_title) and '(' in o['nom']:
                    best_title = o['nom']
                    
            # Meilleure description
            best_desc = primary['desc']
            for o in others:
                if ('[Méthode]' in o['desc']) and ('[Méthode]' not in best_desc or len(o['desc']) > len(best_desc)):
                    best_desc = o['desc']
                    
            # Consolidation de TOUS les liens uniques
            unique_sources = []
            seen_keys = set()
            for item in [primary] + others:
                refs = [x.strip() for x in item['pdf_ref'].split('|') if x.strip()]
                targets = [x.strip() for x in item['target'].split('|') if x.strip()]
                for idx, ref in enumerate(refs):
                    t = targets[idx] if idx < len(targets) else (targets[0] if targets else '')
                    fn = parse_pdf_filename(ref, t)
                    pg = parse_page_number(ref, t)
                    k = (fn.lower(), pg)
                    if k not in seen_keys:
                        seen_keys.add(k)
                        unique_sources.append({
                            'ref': ref,
                            'target': t or f"Sources/sources_lisibles/{fn}#page={pg}",
                            'filename': fn,
                            'page': pg
                        })
                if not refs and item['target']:
                    t = item['target']
                    fn = parse_pdf_filename('', t)
                    pg = parse_page_number('', t)
                    k = (fn.lower(), pg)
                    if k not in seen_keys:
                        seen_keys.add(k)
                        unique_sources.append({
                            'ref': f"{fn}, p. {pg}",
                            'target': t,
                            'filename': fn,
                            'page': pg
                        })
                        
            combined_refs = " | ".join(s['ref'] for s in unique_sources)
            combined_targets = " | ".join(s['target'] for s in unique_sources)
            primary_target = unique_sources[0]['target'] if unique_sources else primary['target']
            
            final_rows.append({
                'id': primary['id'],
                'partie': primary['partie'],
                'chapitre': primary['chapitre'],
                'section': primary['section'],
                'nom': best_title,
                'desc': best_desc,
                'pdf_ref': combined_refs,
                'target': combined_targets,
                'primary_target': primary_target
            })
            
    # 3. Réécriture ultra-rapide de la feuille (remplace feuille en mémoire)
    print(f"⚡ Réécriture rapide des {len(final_rows)} lignes consolidées...")
    sheet_title = sheet.title
    wb.remove(sheet)
    new_sheet = wb.create_sheet(title=sheet_title)
    
    # Headers
    for c, h in enumerate(headers, start=1):
        new_sheet.cell(1, c).value = h
        
    for r_idx, row in enumerate(final_rows, start=2):
        new_sheet.cell(r_idx, 1).value = row['id']
        new_sheet.cell(r_idx, 2).value = row['partie']
        new_sheet.cell(r_idx, 3).value = row['chapitre']
        new_sheet.cell(r_idx, 4).value = row['section']
        new_sheet.cell(r_idx, 5).value = row['nom']
        new_sheet.cell(r_idx, 6).value = row['desc']
        new_sheet.cell(r_idx, 7).value = row['pdf_ref']
        
        target_cell = new_sheet.cell(r_idx, 8)
        target_cell.value = row['target']
        if row['primary_target']:
            target_cell.hyperlink = Hyperlink(ref=f"H{r_idx}", target=row['primary_target'])
            
    backup_excel(file_path)
    wb.save(file_path)
    dt = time.time() - t0
    print(f"💾 Sauvegarde réussie en {dt:.1f}s : {file_name}")
    print(f"✨ Nouvelles lignes totales : {len(final_rows)} techniques uniques (au lieu de {len(rows_data)}).")
    return len(rows_data), len(final_rows)

def run_all():
    print("=" * 60)
    print("🚀 LANCEMENT DU PROTOCOLE ULTRA-RAPIDE DE DÉDUPLICATION DES EXCEL")
    print("=" * 60)
    
    total_before = 0
    total_after = 0
    
    for title, fname in [
        ('Techniques Cartes', 'Livre Techniques Cartes.xlsx'),
        ('Techniques Pièces', 'Livre Techniques Pièces.xlsx'),
        ('Tours de Cartes', 'Livre Tours de Cartes.xlsx')
    ]:
        b, a = deduplicate_workbook_fast(title, fname)
        total_before += b
        total_after += a
        
    print("\n" + "=" * 60)
    print(f"🎉 BILAN GLOBAL DE DÉDUPLICATION :")
    print(f"   • Avant : {total_before:,} lignes")
    print(f"   • Après : {total_after:,} lignes uniques consolidées")
    print(f"   • Lignes doublons supprimées : {total_before - total_after:,}")
    print("=" * 60)

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description="Outil de déduplication et consolidation des techniques Excel")
    parser.add_argument('--book', choices=['cartes', 'pieces', 'tours'], default=None, help="Livre spécifique à dédupliquer")
    parser.add_argument('--all', action='store_true', help="Dédupliquer les 3 livres")
    args = parser.parse_args()

    if args.book:
        config = BOOKS_CONFIG[args.book]
        deduplicate_workbook_fast(config['title'], os.path.basename(config['excel']))
    else:
        run_all()
