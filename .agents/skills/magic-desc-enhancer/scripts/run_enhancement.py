import os
import sys
import argparse
import openpyxl

# Ajout du module partagé au path
_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_PROJECT_ROOT = os.path.abspath(os.path.join(_SCRIPT_DIR, '..', '..', '..', '..'))
if _PROJECT_ROOT not in sys.path:
    sys.path.insert(0, _PROJECT_ROOT)

from shared.config import (
    ensure_utf8, get_book_config, normalize_book_key,
    find_data_sheet, backup_excel, BOOKS_CONFIG
)

ensure_utf8()

from desc_enhancer import MagicDescEnhancer


def enhance_book(book_key: str, limit: int = None, dry_run: bool = False, partie_filter: str = None, chapitre_filter: str = None) -> int:
    config = get_book_config(book_key)
    normalized_key = config['id']
    excel_path = config['excel']

    if not os.path.exists(excel_path):
        print(f"❌ Fichier introuvable: {excel_path}")
        return 0

    enhancer = MagicDescEnhancer()

    print(f"\n=======================================================")
    print(f"✍️ SKILL AGENT 2 : magic-desc-enhancer")
    print(f"=======================================================")
    print(f"📖 Livre cible     : {config['title']} ({normalized_key})")
    print(f"📊 Fichier Excel   : {os.path.basename(excel_path)}")
    print(f"📑 Limite fiches   : {limit if limit else 'Toutes les fiches'}")
    print(f"🔒 Mode simulation : {'OUI (Dry Run)' if dry_run else 'NON (Écriture Excel)'}\n")

    wb = openpyxl.load_workbook(excel_path)
    sheet = find_data_sheet(wb)

    total_modified = 0
    total_titles_modified = 0
    total_desc_modified = 0
    total_processed = 0

    max_rows = sheet.max_row
    for r in range(2, max_rows + 1):
        if limit and total_processed >= limit:
            break

        tech_id = sheet.cell(r, 1).value
        if not tech_id:
            continue

        total_processed += 1

        partie = str(sheet.cell(r, 2).value or 'PARTIE 1').strip()
        chapitre = str(sheet.cell(r, 3).value or 'Chapitre 1 : Général').strip()
        section = str(sheet.cell(r, 4).value or 'Général').strip()
        nom = str(sheet.cell(r, 5).value or 'Sans titre').strip()

        # Filtrage optionnel par partie ou chapitre
        if partie_filter and partie_filter.lower() not in partie.lower():
            continue
        if chapitre_filter and chapitre_filter.lower() not in chapitre.lower():
            continue

        current_desc = str(sheet.cell(r, 6).value or '').strip()

        # 1. Clarification du Titre (Colonne 5)
        new_nom = enhancer.enhance_title(nom, normalized_key, partie, chapitre, section, current_desc)
        title_changed = (new_nom != nom)
        if title_changed:
            total_titles_modified += 1
            if not dry_run:
                sheet.cell(r, 5).value = new_nom

        # 2. Enrichissement de la Description (Colonne 6)
        desc_changed = False
        has_valid_desc = ('[Effet]' in current_desc and '[Méthode]' in current_desc and len(current_desc) >= 30)
        if not has_valid_desc:
            new_desc = enhancer.enhance_description(normalized_key, partie, chapitre, section, new_nom, current_desc)
            if new_desc != current_desc:
                desc_changed = True
                total_desc_modified += 1
                if not dry_run:
                    sheet.cell(r, 6).value = new_desc

        if title_changed or desc_changed:
            total_modified += 1
            if total_modified <= 10 or total_modified % 1000 == 0:
                print(f"   [{tech_id}] Ligne {r}")
                if title_changed:
                    print(f"      Titre : '{nom}' ➔ '{new_nom}'")
                if desc_changed:
                    print(f"      Desc  : {new_desc[:90]}...")

    if not dry_run and total_modified > 0:
        backup_excel(excel_path)
        wb.save(excel_path)
        print(f"\n💾 Sauvegarde réussie dans {os.path.basename(excel_path)} :")
        print(f"   - {total_titles_modified:,} titres clarifiés")
        print(f"   - {total_desc_modified:,} descriptions enrichies")
        print(f"   - {total_modified:,} lignes modifiées sur {total_processed:,} fiches traitées.")
    else:
        print(f"\n🏁 Fin de l'opération ({total_modified} fiches modifiées sur {total_processed} traitées).")

    return total_modified

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description="Skill Agent 2 : magic-desc-enhancer")
    parser.add_argument('--book', choices=['cartes', 'pieces', 'tours', 'cartes_techniques', 'pieces_techniques', 'cartes_tours'], default='cartes', help="Clé du livre")
    parser.add_argument('--all', action='store_true', help="Traiter les 3 livres")
    parser.add_argument('--partie', type=str, default=None, help="Filtrer sur une partie spécifique")
    parser.add_argument('--chapitre', type=str, default=None, help="Filtrer sur un chapitre spécifique")
    parser.add_argument('--limit', type=int, default=None, help="Nombre max de fiches à enrichir (test par lot)")
    parser.add_argument('--dry-run', action='store_true', help="Mode simulation")
    args = parser.parse_args()

    total_enhanced = 0
    if args.all:
        for bk in BOOKS_CONFIG.keys():
            total_enhanced += enhance_book(bk, limit=args.limit, dry_run=args.dry_run, partie_filter=args.partie, chapitre_filter=args.chapitre)
    else:
        total_enhanced = enhance_book(args.book, limit=args.limit, dry_run=args.dry_run, partie_filter=args.partie, chapitre_filter=args.chapitre)

    print(f"\n🏁 Total enrichissements : {total_enhanced} fiches modifiées.")
