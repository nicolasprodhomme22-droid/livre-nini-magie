import openpyxl
import os
import sys
import re

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_PROJECT_ROOT = os.path.abspath(os.path.join(_SCRIPT_DIR, '..', '..', '..', '..'))
if _PROJECT_ROOT not in sys.path:
    sys.path.insert(0, _PROJECT_ROOT)

from shared.config import find_data_sheet

EXPECTED_HEADERS = [
    "Clé d'identification",
    "Partie",
    "Chapitre",
    "Section",
    "Nom technique / tour",
    "Description (Quoi et Comment)",
    "PDF concerné avec numéro de page",
    "Lien vers la page du pdf"
]

def sanitize_text(val: str) -> str:
    """Neutralise les caractères d'attaque ou de formule Excel et nettoie les caractères non imprimables."""
    if not val:
        return ""
    s = str(val).strip()
    while s and s[0] in ('=', '+', '-', '@'):
        s = s[1:].strip()
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', s)

class ExcelManager:
    def __init__(self, excel_path: str, id_prefix: str):
        self.excel_path = os.path.abspath(excel_path)
        self.id_prefix = id_prefix
        self.wb = openpyxl.load_workbook(self.excel_path)
        
        # Trouver l'onglet de données via le module partagé
        self.sheet = find_data_sheet(self.wb)

        self._ensure_headers()
        self.current_max_id = self._find_max_id()
        self.initial_max_id = self.current_max_id
        self.updated_count = 0
        self.added_count = 0
        self._load_existing_techniques()

    def _ensure_headers(self):
        """Vérifie ou initialise les 8 en-têtes officiels."""
        current_headers = [self.sheet.cell(1, c).value for c in range(1, 9)]
        if not any(current_headers):
            for c, h in enumerate(EXPECTED_HEADERS, start=1):
                self.sheet.cell(1, c).value = h

    def _find_max_id(self) -> int:
        """Trouve le numéro d'ID maximum actuel pour incrémenter de façon unique."""
        max_num = 0
        pattern = re.compile(rf"{re.escape(self.id_prefix)}_(\d+)")
        for r in range(2, self.sheet.max_row + 1):
            val = str(self.sheet.cell(r, 1).value or '')
            match = pattern.search(val)
            if match:
                num = int(match.group(1))
                if num > max_num:
                    max_num = num
        return max_num

    def _load_existing_techniques(self):
        """Indexe les techniques déjà existantes pour éviter strictement les doublons et mettre à jour les liens."""
        self.existing_names = {}
        self.existing_tuples = set()
        for r in range(2, self.sheet.max_row + 1):
            nom = str(self.sheet.cell(r, 5).value or '').strip().lower()
            ref = str(self.sheet.cell(r, 7).value or '').strip().lower()
            if nom:
                if nom not in self.existing_names:
                    self.existing_names[nom] = r
                self.existing_tuples.add((nom, ref))

    def is_duplicate(self, nom: str, pdf_reference: str = "") -> bool:
        """Détecte si une technique est déjà présente. Met à jour la référence si c'est un nouveau PDF."""
        nom_clean = nom.strip().lower()
        if not nom_clean:
            return True
            
        is_dup = False
        if nom_clean in self.existing_names:
            is_dup = True
            
        if pdf_reference:
            ref_clean = pdf_reference.strip().lower()
            if (nom_clean, ref_clean) in self.existing_tuples:
                return True
                
            if is_dup:
                # C'est un doublon de nom, mais avec une nouvelle référence.
                # On ajoute la référence dans la colonne 7 pour le link_builder.
                row_idx = self.existing_names[nom_clean]
                cell = self.sheet.cell(row=row_idx, column=7)
                current_ref = str(cell.value or '').strip()
                if current_ref:
                    # On vérifie grossièrement que la référence n'y est pas déjà
                    if pdf_reference.strip() not in current_ref:
                        cell.value = current_ref + " | " + pdf_reference.strip()
                        self.updated_count += 1
                else:
                    cell.value = pdf_reference.strip()
                    self.updated_count += 1
                    
                self.existing_tuples.add((nom_clean, ref_clean))
                
        return is_dup

    def add_technique(self, partie: str, chapitre: str, section: str, nom: str,
                      description: str, pdf_reference: str, pdf_link: str) -> str:
        """Ajoute une nouvelle technique respectant strictement les 8 colonnes."""
        self.current_max_id += 1
        new_id = f"{self.id_prefix}_{self.current_max_id:03d}"
        
        # Règle stricte pour la section : jamais vide, "Général" par défaut
        sec_val = section.strip() if section and section.strip() else "Général"
        
        next_row = self.sheet.max_row + 1
        self.sheet.cell(next_row, 1).value = new_id
        self.sheet.cell(next_row, 2).value = sanitize_text(partie)
        self.sheet.cell(next_row, 3).value = sanitize_text(chapitre)
        self.sheet.cell(next_row, 4).value = sanitize_text(sec_val)
        self.sheet.cell(next_row, 5).value = sanitize_text(nom)
        self.sheet.cell(next_row, 6).value = sanitize_text(description)
        self.sheet.cell(next_row, 7).value = sanitize_text(pdf_reference)
        
        link_cell = self.sheet.cell(next_row, 8)
        if pdf_link:
            # Extraction du numéro de page pour le libellé
            page_m = re.search(r'#page=(\d+)', pdf_link)
            label = f"Ouvrir page {page_m.group(1)}" if page_m else "Ouvrir la page"
            link_cell.value = label
            link_cell.hyperlink = pdf_link
        else:
            link_cell.value = ""

        # Enregistrement pour déduplication immédiate
        nom_clean = nom.strip().lower()
        self.existing_names[nom_clean] = next_row
        if pdf_reference:
            self.existing_tuples.add((nom_clean, pdf_reference.strip().lower()))

        self.added_count += 1
        return new_id

    @property
    def has_changes(self) -> bool:
        """Indique si des ajouts ou des enrichissements de liens ont été effectués."""
        return self.added_count > 0 or self.updated_count > 0

    def save(self):
        """Sauvegarde les modifications dans le fichier Excel."""
        self.wb.save(self.excel_path)
